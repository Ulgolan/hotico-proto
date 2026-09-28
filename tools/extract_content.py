#!/usr/bin/env python3
"""HOTICO content extractor — lap C-1b.

Usage:  python tools/extract_content.py <lang>      (lang = fr | en | ro)

Reads  content/source/hotico-content-<lang>.xlsx  (openpyxl, read-only load;
the workbook is never written) and writes
  content/<lang>.json         — the content bible, schema = key tree of the
                                C-1 content/fr.json (page -> section -> element)
  content/<lang>-review.md    — the same content rendered for human review

Laws this script obeys:
  - Cells are located by what they ARE (their Romanian scaffold label, their
    URL shape, their column header, their place in a block) — never by fixed
    coordinates. Rows and columns drift between the three workbooks.
  - Every string is stored verbatim: no trim, no numbering strip, no quote or
    space cleanup. Line breaks are stored as \\n.
  - A slot with no content is null. Nothing is translated or borrowed.
  - Any cell that cannot be placed with confidence goes to the top-level
    "_unplaced" list as {sheet, cell, text}. Nothing is forced.
Output is deterministic: same workbook bytes in, same bytes out.
"""

import hashlib
import json
import os
import re
import sys

import openpyxl
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
LANGS = ("fr", "en", "ro")
LAP = "C-1b"
LAW = "verbatim — copy is reproduced exactly as stored in the spreadsheet"

# sheet title (normalised) -> page key
SHEET_KEYS = {
    "homepage": "home",
    "s.areola": "areola",
    "s.alopecie": "alopecie",
    "s.cicatrici": "cicatrici",
    "spr": "spr",
    "eyeliner": "eyeliner",
    "buze": "buze",
    "lppara": "lp_para",
    "lpcosmetic": "lp_cosmetic",
    "cont": "cont",
}
SERVICE_PAGES = ("areola", "alopecie", "cicatrici", "spr", "eyeliner", "buze")
PAGE_ORDER = ("home",) + SERVICE_PAGES + ("lp_para", "lp_cosmetic", "cont")


# ---------------------------------------------------------------- matching --
# Normalisation is used ONLY to recognise a cell. Stored values are raw.

def norm(s):
    return " ".join(s.split()).lower()


def is_youtube(s):
    return "youtu.be/" in s or "youtube.com/" in s


def is_vimeo(s):
    return "vimeo.com/" in s


def is_url(s):
    return is_youtube(s) or is_vimeo(s)


def is_title_label(s):
    return norm(s) == "titlu video"


def is_cta_label(s):
    return norm(s).startswith("cta sub titlu")


def is_questions_label(s):
    return norm(s).startswith("cta intrebari")


def is_tabs_label(s):
    return "mini meniu" in norm(s)


def is_step2_label(s):
    return norm(s).startswith("in formularul de programare la pasul 2")


def is_step3_label(s):
    return norm(s).startswith("pasul 3")


def is_branching_header(s):
    return norm(s).startswith("ce se intampla daca bifeaza")


def is_column_header(s):
    n = norm(s)
    return (n.startswith("link ") or n.startswith("descriere raspuns")
            or n.startswith("decriere raspuns"))


def is_annotation(s):
    # "<- TITLU", "<- TEXT EXPLICATIV": arrows pointing at a neighbour cell
    return s.strip().startswith("<-")


def is_form_note(s):
    # Romanian builder instructions inside the medical-form block
    n = norm(s)
    return n.startswith("daca bi") or n.startswith("trebuie casuta")


def is_label(s):
    return (is_title_label(s) or is_cta_label(s) or is_questions_label(s)
            or is_tabs_label(s) or is_step2_label(s) or is_step3_label(s)
            or is_branching_header(s) or is_column_header(s)
            or is_annotation(s))


PREFIX_RE = re.compile(r"^([A-Za-zăâîșțĂÂÎȘȚ]+): ?(.*)$", re.DOTALL)


def split_prefix(s):
    """'titlu: X' -> ('titlu', 'X'). Only one separator space is consumed."""
    m = PREFIX_RE.match(s)
    if not m or m.group(1).lower() in ("http", "https"):
        return None, s
    return m.group(1), m.group(2)


# ------------------------------------------------------------------ sheet --

class Sheet:
    """A sheet's non-empty cells plus a ledger of which ones were placed."""

    def __init__(self, ws):
        self.title = ws.title
        self.cells = {}
        for r, row in enumerate(ws.iter_rows(values_only=True), 1):
            for c, v in enumerate(row, 1):
                if v is None:
                    continue
                if not isinstance(v, str):
                    v = str(v)
                self.cells[(r, c)] = v.replace("\r\n", "\n").replace("\r", "\n")
        self.used = set()

    def take(self, rc):
        if rc is None:
            return None
        self.used.add(rc)
        return self.cells[rc]

    def find(self, pred, unused_only=True, min_row=1, max_row=None):
        out = []
        for rc in sorted(self.cells):
            if unused_only and rc in self.used:
                continue
            if rc[0] < min_row or (max_row is not None and rc[0] > max_row):
                continue
            if pred(self.cells[rc]):
                out.append(rc)
        return out

    def first(self, pred, **kw):
        hits = self.find(pred, **kw)
        return hits[0] if hits else None

    def row(self, r, unused_only=True):
        return [(r, c) for (rr, c) in sorted(self.cells)
                if rr == r and (not unused_only or (rr, c) not in self.used)]

    def rows(self):
        return sorted({r for (r, _) in self.cells})

    def unplaced(self):
        return [{"sheet": self.title,
                 "cell": f"{get_column_letter(c)}{r}",
                 "text": self.cells[(r, c)]}
                for (r, c) in sorted(self.cells) if (r, c) not in self.used]


def coord(rc):
    return f"{get_column_letter(rc[1])}{rc[0]}"


def sheet_labels_row1(sh):
    """Row 1 verbatim, keyed by cell coordinate (internal notes stay put)."""
    out = {}
    for rc in sh.row(1, unused_only=False):
        out[coord(rc)] = sh.take(rc)
    return out or None


def texts_right_of(sh, rc, pred=lambda s: True):
    """Unused non-URL text cells in rc's row, right of rc."""
    return [x for x in sh.row(rc[0])
            if x[1] > rc[1] and not is_url(sh.cells[x]) and pred(sh.cells[x])]


def labelled_value(sh, label_pred):
    """(label_text, value_text) for a 'label | value' row. Either may be None."""
    lab = sh.first(label_pred, unused_only=False)
    if lab is None:
        return None, None
    label = sh.take(lab)
    vals = texts_right_of(sh, lab, lambda s: not is_label(s))
    return label, sh.take(vals[0]) if vals else None


# ------------------------------------------------------------------- home --

def extract_home(sh):
    labels = sheet_labels_row1(sh)
    col_role = {}
    for rc in sh.row(1, unused_only=False):
        n = norm(sh.cells[rc])
        if n == "denumire":
            col_role[rc[1]] = "title"
        elif n.startswith("descriere cta"):
            col_role[rc[1]] = "cta"
        elif n.startswith("descriere text"):
            col_role[rc[1]] = "body"

    # carousel: every row carrying a video URL
    carousel = []
    for r in sh.rows():
        if r == 1:
            continue
        cells = sh.row(r)
        if not any(is_url(sh.cells[x]) for x in cells):
            continue
        item = {"title": None, "cta": None, "body": None,
                "youtube_url": None, "vimeo_url": None}
        for x in cells:
            v = sh.cells[x]
            if is_youtube(v) and item["youtube_url"] is None:
                item["youtube_url"] = sh.take(x)
            elif is_vimeo(v) and item["vimeo_url"] is None:
                item["vimeo_url"] = sh.take(x)
            elif not is_url(v) and x[1] in col_role and item[col_role[x[1]]] is None:
                item[col_role[x[1]]] = sh.take(x)
        carousel.append({k: item[k] for k in
                         ("title", "cta", "body", "youtube_url", "vimeo_url")})

    # hero
    hero_lab = sh.first(lambda s: norm(s).startswith("sub logo"))
    hero_texts = texts_right_of(sh, hero_lab) if hero_lab else []
    hero_row_label = sh.take(hero_lab)
    tagline = sh.take(hero_texts[0]) if len(hero_texts) > 0 else None
    intro = sh.take(hero_texts[1]) if len(hero_texts) > 1 else None
    cta_row_label, cta = labelled_value(
        sh, lambda s: norm(s).startswith("deasupra de video"))

    # services menu: the cell that lists the service names
    svc = sh.first(lambda s: "numele serviciilor" in norm(s))
    svc_row_label = None
    if svc:
        others = [x for x in sh.row(svc[0]) if x != svc]
        svc_row_label = sh.take(others[0]) if others else None
    svc_text = sh.take(svc)

    # form block anchored on the GDPR checkbox row
    gdpr_lab = sh.first(lambda s: norm(s).startswith("bifa gdpr"))
    canc_lab = sh.first(lambda s: norm(s).startswith("bifa anularea"))
    button = sh.first(lambda s: norm(s).startswith("buton"))

    def pair(lab):
        if lab is None:
            return None
        vals = texts_right_of(sh, lab)
        return {"row_label": sh.take(lab),
                "text": sh.take(vals[0]) if vals else None}

    form_heading = form_cta = None
    heading_row = None
    if gdpr_lab:
        above = [r for r in sh.rows() if r < gdpr_lab[0] and sh.row(r)]
        if above:
            heading_row = above[-1]
            hc = sh.row(heading_row)
            form_heading = sh.take(hc[0])
            form_cta = sh.take(hc[1]) if len(hc) > 1 else None
    gdpr = pair(gdpr_lab)
    cancellation = pair(canc_lab)
    next_button = sh.take(button)

    # steps: between the services row and the form heading row
    pasii_heading = None
    steps = []
    lo = svc[0] if svc else 1
    hi = heading_row if heading_row else (max(sh.rows()) + 1)
    region = [r for r in sh.rows() if lo < r < hi and sh.row(r)]
    if region:
        hc = sh.row(region[0])
        if len(hc) == 1:
            pasii_heading = sh.take(hc[0])
            region = region[1:]
    for r in region:
        rc = sh.row(r)
        if len(rc) == 2:
            steps.append({"title": sh.take(rc[0]), "body": None,
                          "items": [sh.take(rc[1])]})
        elif len(rc) == 1 and steps:
            steps[-1]["items"].append(sh.take(rc[0]))
        # anything else stays unplaced

    # reviews heading: the first text after the button row
    reviews_heading = None
    if button:
        after = [x for r in sh.rows() if r > button[0] for x in sh.row(r)]
        if after:
            reviews_heading = sh.take(after[0])

    return {
        "hero": {
            "row_label": hero_row_label,
            "tagline": tagline,
            "h1": None,
            "intro": intro,
            "cta": cta,
            "cta_row_label": cta_row_label,
        },
        "services_menu": {"row_label": svc_row_label, "text": svc_text},
        "carousel": carousel or None,
        "pasii": {
            "heading": pasii_heading,
            "intro": None,
            "steps": steps or None,
        },
        "form": {
            "heading": form_heading,
            "cta": form_cta,
            "gdpr": gdpr,
            "cancellation": cancellation,
            "next_button": next_button,
            "labels": None,
            "options": None,
            "confirmation": None,
        },
        "reviews": {"heading": reviews_heading, "items": None},
        "footer": None,
        "_sheet_labels": labels,
    }


# ---------------------------------------------------------------- service --

def extract_service(sh):
    labels = sheet_labels_row1(sh)
    title_lab = sh.first(is_title_label, unused_only=False)
    cta_lab = sh.first(is_cta_label, unused_only=False)
    q_lab = sh.first(is_questions_label, unused_only=False)

    # title: first plain text from the label onwards, before the CTA row
    title = None
    if title_lab:
        stop = cta_lab[0] if cta_lab else title_lab[0] + 1
        for rc in sorted(sh.cells):
            if not (title_lab[0] <= rc[0] < stop):
                continue
            if rc[0] == title_lab[0] and rc[1] <= title_lab[1]:
                continue
            v = sh.cells[rc]
            if is_url(v) or is_label(v):
                continue
            sh.used.add(rc)
            title = v
            break
    title_row_label = sh.take(title_lab)

    cta_row_label, cta = labelled_value(sh, is_cta_label)
    q_row_label, q_intro = labelled_value(sh, is_questions_label)

    # tabs: label cell, items stacked below it in the same column
    tabs = None
    tl = sh.first(is_tabs_label, unused_only=False)
    if tl:
        items = []
        r = tl[0] + 1
        while (r, tl[1]) in sh.cells:
            items.append(sh.take((r, tl[1])))
            r += 1
        tabs = {"row_label": sh.take(tl), "items": items or None}

    step2_lab = sh.first(is_step2_label, unused_only=False)
    step3_lab = sh.first(is_step3_label, unused_only=False)

    # FAQ: rows between the questions label and the step-2 form row
    faq = []
    lo = q_lab[0] if q_lab else (cta_lab[0] if cta_lab else 1)
    ends = [x[0] for x in (step2_lab, step3_lab) if x]
    hi = min(ends) if ends else max(sh.rows()) + 1
    for r in sh.rows():
        if not (lo < r < hi):
            continue
        cells = sh.row(r)
        if not cells:
            continue
        yt = [x for x in cells if is_youtube(sh.cells[x])]
        vm = [x for x in cells if is_vimeo(sh.cells[x])]
        tx = [x for x in cells if not is_url(sh.cells[x])
              and not is_label(sh.cells[x])]
        faq.append({
            "question": sh.take(tx[0]) if len(tx) > 0 else None,
            "answer": sh.take(tx[1]) if len(tx) > 1 else None,
            "youtube_url": sh.take(yt[0]) if yt else None,
            "vimeo_url": sh.take(vm[0]) if vm else None,
        })

    return {
        "intro": {
            "title": title,
            "title_row_label": title_row_label,
            "cta": cta,
            "cta_row_label": cta_row_label,
            "questions_intro": q_intro,
            "questions_row_label": q_row_label,
        },
        "_sheet_labels": labels,
        "tabs": tabs,
        "faq": faq or None,
        "pricing": None,
        "gallery": None,
        "form_notes": extract_form_notes(sh, step2_lab, step3_lab),
    }


def extract_form_notes(sh, step2_lab, step3_lab):
    if step2_lab is None and step3_lab is None:
        return None

    step2 = None
    if step2_lab:
        vals = texts_right_of(sh, step2_lab, lambda s: not is_label(s))
        step2 = {"row_label": sh.take(step2_lab),
                 "text": sh.take(vals[0]) if vals else None}

    step3 = intro = branching_header = questions = free_text = None
    if step3_lab:
        vals = texts_right_of(sh, step3_lab, lambda s: not is_label(s))
        step3 = {"row_label": sh.take(step3_lab),
                 "heading": sh.take(vals[0]) if vals else None}

        bh = sh.first(is_branching_header, unused_only=False,
                      min_row=step3_lab[0])
        branching_col = bh[1] if bh else None
        branching_header = sh.take(bh)

        def plain(rc):
            v = sh.cells[rc]
            return not is_label(v) and not is_url(v)

        # intro: first row after the step-3 row carrying plain text
        intro_row = step3_lab[0]
        for r in sh.rows():
            if r <= step3_lab[0]:
                continue
            cand = [x for x in sh.row(r) if plain(x)]
            if cand:
                intro = sh.take(cand[0])
                intro_row = r
                break

        # questions: the first contiguous run of rows after the intro
        # (blank rows before the run are skipped; a blank row ends it)
        qrows = []
        for r in range(intro_row + 1, max(sh.rows()) + 1):
            if sh.row(r):
                qrows.append(r)
            elif qrows:
                break

        qlist = []
        for qr in qrows:
            cells = sh.row(qr)
            branching = None
            if branching_col is not None:
                bc = [x for x in cells if x[1] == branching_col]
                if bc:
                    branching = sh.take(bc[0])
            rest = [x for x in sh.row(qr) if plain(x)]
            if not rest:
                continue
            q = sh.take(rest[0])
            notes = [x for x in rest[1:] if is_form_note(sh.cells[x])]
            others = [x for x in rest[1:] if not is_form_note(sh.cells[x])]
            qlist.append({
                "row": qr,
                "question": q,
                "note": sh.take(notes[0]) if notes else None,
                "options": sh.take(others[0]) if len(others) > 0 else None,
                "extra": sh.take(others[1]) if len(others) > 1 else None,
                "branching": branching,
            })
        questions = qlist or None

        # free text: next non-empty row after the question run
        if qrows:
            after = [r for r in sh.rows() if r > qrows[-1] and sh.row(r)]
            if after:
                cells = [x for x in sh.row(after[0]) if plain(x)]
                if cells:
                    prompt = sh.take(cells[0])
                    notes = [x for x in cells[1:] if is_form_note(sh.cells[x])]
                    free_text = {"prompt": prompt,
                                 "note": sh.take(notes[0]) if notes else None}

    return {
        "step2": step2,
        "step3": step3,
        "intro": intro,
        "branching_column_header": branching_header,
        "questions": questions,
        "free_text": free_text,
    }


# --------------------------------------------------------------- landings --

def extract_landing(sh):
    yt_h = sh.first(lambda s: norm(s).startswith("link youtube")
                    or norm(s).startswith("link yt"))
    vm_h = sh.first(lambda s: norm(s).startswith("link vimeo"))
    yt = sh.first(is_youtube)
    vm = sh.first(is_vimeo)
    return {
        "youtube_url": sh.take(yt),
        "vimeo_url": sh.take(vm),
        "text": None,
        "_headers": {"youtube": sh.take(yt_h), "vimeo": sh.take(vm_h)},
    }


def extract_cont(sh):
    vimeo_label = sh.first(lambda s: norm(s) == "link vimeo")
    videos = []
    for yrc in sh.find(is_youtube):
        slot, yurl = split_prefix(sh.cells[yrc])
        row = sh.row(yrc[0])
        t = [x for x in row if x != yrc and
             (split_prefix(sh.cells[x])[0] or "").lower() == "titlu"]
        c = [x for x in row if x != yrc and
             (split_prefix(sh.cells[x])[0] or "").lower() == "cta"]
        v = [x for x in sh.find(is_vimeo)
             if slot is not None and split_prefix(sh.cells[x])[0] == slot]
        verb = {
            "youtube": sh.take(yrc),
            "title": sh.take(t[0]) if t else None,
            "cta": sh.take(c[0]) if c else None,
            "vimeo": sh.take(v[0]) if v else None,
        }
        videos.append({
            "slot": slot,
            "title": split_prefix(verb["title"])[1] if verb["title"] else None,
            "cta": split_prefix(verb["cta"])[1] if verb["cta"] else None,
            "youtube_url": yurl,
            "vimeo_url": split_prefix(verb["vimeo"])[1] if verb["vimeo"] else None,
            "_verbatim": verb,
        })
    return {"vimeo_row_label": sh.take(vimeo_label), "videos": videos or None}


# ------------------------------------------------------------------ build --

def sha256(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(65536), b""):
            h.update(chunk)
    return h.hexdigest()


def extract(lang):
    rel = f"content/source/hotico-content-{lang}.xlsx"
    path = os.path.join(ROOT, rel)
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = {"_meta": {
        "language": lang,
        "source": rel,
        "source_sha256": sha256(path),
        "lap": LAP,
        "law": LAW,
        "sheets": [ws.title for ws in wb.worksheets],
    }}
    pages = {}
    unplaced = []
    for ws in wb.worksheets:
        sh = Sheet(ws)
        key = SHEET_KEYS.get(norm(ws.title).replace(" ", ""))
        if key == "home":
            pages[key] = extract_home(sh)
        elif key in SERVICE_PAGES:
            pages[key] = extract_service(sh)
        elif key in ("lp_para", "lp_cosmetic"):
            pages[key] = extract_landing(sh)
        elif key == "cont":
            pages[key] = extract_cont(sh)
        unplaced.extend(sh.unplaced())
    wb.close()
    for key in PAGE_ORDER:
        out[key] = pages.get(key)
    out["_unplaced"] = unplaced
    return out


# ----------------------------------------------------------------- review --

NULL = "_(null)_"

PAGE_NAMES = {
    "fr": {"areola": "Aréole", "alopecie": "Alopécie", "cicatrici": "Cicatrices",
           "spr": "Sourcils", "eyeliner": "Eyeliner", "buze": "Lèvres",
           "lp_para": "LP paramédical", "lp_cosmetic": "LP cosmétique",
           "cont": "Contenu / guides"},
    "en": {"areola": "Areola", "alopecie": "Alopecia", "cicatrici": "Scars",
           "spr": "Brows", "eyeliner": "Eyeliner", "buze": "Lips",
           "lp_para": "LP paramedical", "lp_cosmetic": "LP cosmetic",
           "cont": "Content / guides"},
    "ro": {"areola": "Areola", "alopecie": "Alopecie", "cicatrici": "Cicatrici",
           "spr": "Sprâncene", "eyeliner": "Eyeliner", "buze": "Buze",
           "lp_para": "LP paramedical", "lp_cosmetic": "LP cosmetic",
           "cont": "Conținut / ghiduri"},
}


def quote(v):
    return "\n".join("> " + ln if ln else ">" for ln in v.split("\n"))


def field(label, v, block=False):
    if v is None:
        return [f"- **{label}:** {NULL}"]
    if block or "\n" in v:
        return [f"- **{label}:**", "", quote(v), ""]
    return [f"- **{label}:** {v}"]


def url_field(label, v):
    # a video slot left empty in the sheet is flagged, not silently nulled
    if v is None:
        return [f"- **{label}:** **— MISSING —**"]
    return field(label, v)


def walk_urls(o, path, pred, acc):
    if isinstance(o, dict):
        for k, v in o.items():
            if k.startswith("_"):
                continue
            walk_urls(v, f"{path}.{k}" if path else k, pred, acc)
    elif isinstance(o, list):
        for i, v in enumerate(o):
            walk_urls(v, f"{path}[{i}]", pred, acc)
    elif isinstance(o, str) and path.endswith("_url") and pred(o):
        acc.append((path, o))


def render_sheet_labels(L, labels):
    L += ["### Sheet column labels (Romanian scaffolding, verbatim)", ""]
    if labels:
        for k, v in labels.items():
            L.append(f"- `{k}` — {v}")
    else:
        L.append(NULL)
    L.append("")


def render_home(L, h):
    L += ["## 1. Home Page", "", "### Hero", ""]
    hero = h["hero"]
    L += field("Row label", hero["row_label"])
    L += field("Tagline", hero["tagline"])
    L += field("H1", hero["h1"])
    L += field("CTA", hero["cta"])
    L += field("CTA row label", hero["cta_row_label"])
    L += field("Intro", hero["intro"])
    L += ["", "### Services menu", ""]
    L += field("Row label", h["services_menu"]["row_label"])
    L += field("Text", h["services_menu"]["text"])
    car = h["carousel"] or []
    L += ["", f"### Carousel ({len(car)} items)", ""]
    for i, c in enumerate(car, 1):
        L += [f"#### Carousel {i} — {c['title']}", ""]
        L += field("Title", c["title"])
        L += field("CTA", c["cta"])
        L += field("YouTube", c["youtube_url"])
        L += field("Vimeo", c["vimeo_url"])
        L += field("Body", c["body"], block=True)
        L += [""]
    if car:
        L += [""]
    p = h["pasii"]
    L += [f"### {p['heading']}", ""]
    L += field("Heading", p["heading"])
    L += field("Intro", p["intro"])
    L += [""]
    for s in p["steps"] or []:
        L += [f"#### Phase — {s['title']}", ""]
        L += field("Title", s["title"])
        L += field("Body", s["body"])
        L += [""]
        for i, it in enumerate(s["items"] or [], 1):
            L += field(f"Item {i}", it)
            L += [""]
        L += [""]
    f = h["form"]
    L += ["### Form", ""]
    L += field("Heading", f["heading"])
    L += field("CTA", f["cta"])
    for key, name in (("gdpr", "GDPR"), ("cancellation", "Cancellation")):
        g = f[key] or {"row_label": None, "text": None}
        L += field(f"{name} row label", g["row_label"])
        L += field(f"{name} text", g["text"])
    L += field("Next button", f["next_button"])
    L += field("Labels", f["labels"])
    L += field("Options", f["options"])
    L += field("Confirmation", f["confirmation"])
    L += ["", "### Reviews", ""]
    L += field("Heading", h["reviews"]["heading"])
    L += field("Items", h["reviews"]["items"])
    L += ["", "### Footer", ""]
    L += field("Footer", h["footer"])
    L += [""]
    render_sheet_labels(L, h["_sheet_labels"])


def render_service(L, n, name, sheet, s):
    L += [f"## {n}. {name}  (sheet `{sheet}`)", "", "### Intro", ""]
    i = s["intro"]
    L += field("Title row label", i["title_row_label"])
    L += field("Title", i["title"])
    L += field("CTA row label", i["cta_row_label"])
    L += field("CTA", i["cta"])
    L += field("Questions row label", i["questions_row_label"])
    L += field("Questions intro", i["questions_intro"])
    L += ["", "### Tabs", ""]
    if s["tabs"]:
        L += field("Row label", s["tabs"]["row_label"])
        for it in s["tabs"]["items"] or []:
            L.append(f"- {it}")
    else:
        L.append(NULL)
    faq = s["faq"] or []
    L += ["", f"### FAQ ({len(faq)} entries)", ""]
    for k, q in enumerate(faq, 1):
        L += [f"#### Q{k}. {q['question']}", ""]
        L += field("Question", q["question"])
        L += url_field("YouTube", q["youtube_url"])
        L += url_field("Vimeo", q["vimeo_url"])
        L += field("Answer", q["answer"], block=True)
        L += [""]
    if faq:
        L += [""]
    L += ["### Pricing", ""]
    L += field("Pricing", s["pricing"])
    L += ["", "### Gallery", ""]
    L += field("Gallery", s["gallery"])
    L += ["", "### Form notes", ""]
    fn = s["form_notes"]
    if fn is None:
        L += [NULL, ""]
    else:
        st2 = fn["step2"] or {"row_label": None, "text": None}
        st3 = fn["step3"] or {"row_label": None, "heading": None}
        L += field("Step 2 row label", st2["row_label"])
        L += field("Step 2 text", st2["text"])
        L += field("Step 3 row label", st3["row_label"])
        L += field("Step 3 heading", st3["heading"])
        L += field("Intro", fn["intro"])
        L += field("Branching column header", fn["branching_column_header"])
        qs = fn["questions"] or []
        L += ["", f"**Questions ({len(qs)}):**", ""]
        for q in qs:
            L += [f"##### Row {q['row']}", ""]
            L += field("Question", q["question"])
            L += field("Note", q["note"])
            L += field("Options", q["options"])
            L += field("Extra", q["extra"])
            L += field("Branching", q["branching"])
            L += [""]
        ft = fn["free_text"] or {"prompt": None, "note": None}
        L += field("Free-text prompt", ft["prompt"])
        L += field("Free-text note", ft["note"])
        L += [""]
    render_sheet_labels(L, s["_sheet_labels"])


def render_landing(L, n, name, sheet, p):
    L += [f"## {n}. {name}  (sheet `{sheet}`)", ""]
    L += field("YouTube", p["youtube_url"])
    L += field("Vimeo", p["vimeo_url"])
    L += field("Text", p["text"])
    L += ["", "**Sheet labels:**", ""]
    for k, v in p["_headers"].items():
        L.append(f"- {k} — {v if v is not None else NULL}")
    L += [""]


def render_cont(L, n, name, sheet, c):
    L += [f"## {n}. {name}  (sheet `{sheet}`)", ""]
    L += field("Vimeo row label", c["vimeo_row_label"])
    L += [""]
    for i, v in enumerate(c["videos"] or [], 1):
        L += [f"### Video {i} — slot `{v['slot']}`", ""]
        L += field("Title", v["title"])
        L += field("CTA", v["cta"])
        L += field("YouTube", v["youtube_url"])
        L += field("Vimeo", v["vimeo_url"])
        L += ["", "**Verbatim cell strings (prefixes intact):**", ""]
        for k, vv in v["_verbatim"].items():
            L.append(f"- `{k}` — {vv if vv is not None else NULL}")
        L += [""]


def render(d):
    lang = d["_meta"]["language"]
    names = PAGE_NAMES[lang]
    sheets = {SHEET_KEYS.get(norm(t).replace(" ", "")): t
              for t in d["_meta"]["sheets"]}
    L = [f"# HOTICO — {lang.upper()} content bible, human review", "",
         f"Rendered verbatim from `{d['_meta']['source']}` via `content/{lang}.json`.",
         "Nothing here is edited, corrected, or normalised. Line breaks inside cells are shown as-is.",
         "`_(null)_` means the spreadsheet had no content for that slot — not that content was dropped.",
         "", "---", ""]
    render_home(L, d["home"])
    n = 2
    for key in PAGE_ORDER[1:]:
        L += ["---", ""]
        p = d.get(key)
        if key in SERVICE_PAGES:
            render_service(L, n, names[key], sheets.get(key), p)
        elif key == "cont":
            render_cont(L, n, names[key], sheets.get(key), p)
        else:
            render_landing(L, n, names[key], sheets.get(key), p)
        n += 1
    for label, pred in (("Vimeo", is_vimeo), ("YouTube", is_youtube)):
        urls = []
        walk_urls({k: v for k, v in d.items() if not k.startswith("_")},
                  "", pred, urls)
        L += ["---", "", f"## {label} URL index", "", "| # | Path | URL |",
              "|---|---|---|"]
        for i, (pth, u) in enumerate(urls, 1):
            L.append(f"| {i} | `{pth}` | {u} |")
        L += ["", f"**Total {label} URLs: {len(urls)}**", ""]
    if "_unplaced" in d:
        L += ["---", "", f"## Unplaced cells ({len(d['_unplaced'])})", "",
              "Cells the extractor could not place with confidence, verbatim.", ""]
        if not d["_unplaced"]:
            L.append("_(none)_")
        for u in d["_unplaced"]:
            L += [f"- `{u['sheet']}` `{u['cell']}` — `{json.dumps(u['text'], ensure_ascii=False)}`"]
        L += [""]
    return "\n".join(L)


# ------------------------------------------------------------------- main --

def main(argv):
    if len(argv) != 2 or argv[1] not in LANGS:
        sys.stderr.write("usage: python tools/extract_content.py <fr|en|ro>\n")
        return 2
    lang = argv[1]
    data = extract(lang)
    with open(os.path.join(ROOT, "content", f"{lang}.json"), "w",
              encoding="utf-8", newline="\n") as f:
        f.write(json.dumps(data, ensure_ascii=False, indent=2) + "\n")
    with open(os.path.join(ROOT, "content", f"{lang}-review.md"), "w",
              encoding="utf-8", newline="\n") as f:
        f.write(render(data))
    print(f"{lang}: wrote content/{lang}.json, content/{lang}-review.md "
          f"({len(data['_unplaced'])} unplaced)")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
